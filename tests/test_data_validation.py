"""Tests for data loading/cleaning and independent checks against the real Excel file."""
import math
from collections import defaultdict

import pandas as pd
import pytest

from src import metrics
from src.data_cleaning import REQUIRED_COLUMNS, MissingColumnsError, clean_data, validate_columns
from src.data_loader import DataLoadError, get_default_data_path, load_raw_data


def make_raw(**overrides):
    """One valid raw row (as read from Excel: dates as text) with optional overrides."""
    row = {
        "Order_ID": "A1",
        "Order_Date": "2024-05-01",
        "Category": "Apparel & Fashion",
        "Product": "Printed T-Shirt",
        "Quantity": 2,
        "Unit_Price_INR": 500.0,
        "Discount_Pct": 0.1,
        "Total_Sales_INR": 900.0,
        "Profit_INR": 200.0,
        "Payment_Method": "UPI",
        "Fulfillment": "Amazon (FBA)",
        "Order_Status": "Delivered",
        "Ship_State": "Karnataka",
    }
    row.update(overrides)
    return row


def raw_frame(*rows):
    return pd.DataFrame(list(rows))


# ---------------------------------------------------------------- columns --
def test_validate_columns_reports_every_missing_column():
    frame = raw_frame(make_raw()).drop(columns=["Profit_INR", "Ship_State"])
    with pytest.raises(MissingColumnsError) as error:
        validate_columns(frame)
    assert sorted(error.value.missing) == ["Profit_INR", "Ship_State"]
    assert "Profit_INR" in str(error.value) and "Ship_State" in str(error.value)


def test_clean_data_rejects_missing_columns():
    with pytest.raises(MissingColumnsError):
        clean_data(raw_frame(make_raw()).drop(columns=["Order_Status"]))


def test_column_names_with_extra_spaces_are_accepted():
    frame = raw_frame(make_raw()).rename(columns={"Order_ID": " Order_ID "})
    cleaned, _ = clean_data(frame)
    assert "Order_ID" in cleaned.columns


# ---------------------------------------------------------------- cleaning --
def test_dates_become_real_datetimes_and_text_is_trimmed():
    cleaned, _ = clean_data(raw_frame(make_raw(Category="  Apparel & Fashion  ", Ship_State=" Karnataka")))
    assert pd.api.types.is_datetime64_any_dtype(cleaned["Order_Date"])
    assert cleaned.loc[0, "Order_Date"] == pd.Timestamp("2024-05-01")
    assert cleaned.loc[0, "Category"] == "Apparel & Fashion"
    assert cleaned.loc[0, "Ship_State"] == "Karnataka"


def test_status_labels_are_made_consistent():
    frame = raw_frame(
        make_raw(Order_ID="A1", Order_Status="delivered "),
        make_raw(Order_ID="A2", Order_Status="CANCELED"),
        make_raw(Order_ID="A3", Order_Status="Returned"),
    )
    cleaned, report = clean_data(frame)
    assert list(cleaned["Order_Status"]) == ["Delivered", "Cancelled", "Returned"]
    assert report["unrecognised_order_statuses"] == []


def test_unrecognised_status_is_reported_not_hidden():
    cleaned, report = clean_data(raw_frame(make_raw(Order_Status="On Hold")))
    assert report["unrecognised_order_statuses"] == ["On Hold"]
    assert len(cleaned) == 1


def test_invalid_dates_and_missing_ids_are_removed_and_counted():
    frame = raw_frame(
        make_raw(Order_ID="A1"),
        make_raw(Order_ID="A2", Order_Date="not a date"),
        make_raw(Order_ID="A3", Order_Date=None),
        make_raw(Order_ID=None),
    )
    cleaned, report = clean_data(frame)
    assert list(cleaned["Order_ID"]) == ["A1"]
    assert report["rows_dropped_invalid_date"] == 2
    assert report["rows_dropped_missing_order_id"] == 1
    assert report["rows_in"] == 4 and report["rows_out"] == 1


def test_exact_duplicates_removed_but_same_id_with_different_content_kept():
    frame = raw_frame(
        make_raw(Order_ID="A1"),
        make_raw(Order_ID="A1"),  # exact duplicate -> removed
        make_raw(Order_ID="A1", Product="Sports Shoes", Total_Sales_INR=1500.0),  # different -> kept
    )
    cleaned, report = clean_data(frame)
    assert report["exact_duplicate_rows_removed"] == 1
    assert report["repeated_order_ids_kept"] == 1
    assert len(cleaned) == 2
    assert metrics.compute_kpis(cleaned)["total_orders"] == 1  # still ONE order


def test_missing_numbers_are_not_invented():
    cleaned, report = clean_data(raw_frame(make_raw(Profit_INR=None)))
    assert math.isnan(cleaned.loc[0, "Profit_INR"])
    assert report["missing_numeric_values_remaining"] == {"Profit_INR": 1}
    assert metrics.compute_kpis(cleaned)["total_profit"] == 0.0


def test_text_numbers_and_negative_values_are_set_to_missing_and_reported():
    frame = raw_frame(
        make_raw(Order_ID="A1", Quantity="abc"),
        make_raw(Order_ID="A2", Total_Sales_INR=-50.0),
        make_raw(Order_ID="A3", Profit_INR=-30.0),  # negative profit = a real loss, kept
    )
    cleaned, report = clean_data(frame)
    assert report["non_numeric_values_set_to_missing"] == {"Quantity": 1}
    assert report["negative_values_set_to_missing"] == {"Total_Sales_INR": 1}
    assert cleaned.loc[2, "Profit_INR"] == -30.0


def test_missing_labels_become_unknown():
    cleaned, report = clean_data(raw_frame(make_raw(Ship_State=None, Payment_Method="  ")))
    assert cleaned.loc[0, "Ship_State"] == "Unknown"
    assert cleaned.loc[0, "Payment_Method"] == "Unknown"
    assert report["missing_text_filled_with_unknown"] == {"Payment_Method": 1, "Ship_State": 1}


def test_cleaning_never_modifies_the_original_frame():
    raw = raw_frame(make_raw(Order_Status=" delivered"), make_raw(Order_ID="A2"))
    before = raw.copy()
    clean_data(raw)
    pd.testing.assert_frame_equal(raw, before)


# ------------------------------------------------------------------ loader --
def test_missing_file_gives_a_clear_error(tmp_path):
    with pytest.raises(DataLoadError) as error:
        load_raw_data(tmp_path / "does_not_exist.xlsx")
    assert "not found" in str(error.value).lower()


def test_unreadable_file_gives_a_clear_error(tmp_path):
    bad = tmp_path / "broken.xlsx"
    bad.write_text("this is not an Excel file")
    with pytest.raises(DataLoadError):
        load_raw_data(bad)


# ------------------------------------------- real dataset, independent check --
DATA_PATH = get_default_data_path()
needs_dataset = pytest.mark.skipif(not DATA_PATH.exists(), reason="dataset file not in data/ folder")


def read_with_openpyxl(path):
    """Read the Excel file WITHOUT pandas, as an independent source of truth."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    header = next(rows)
    index = {name: position for position, name in enumerate(header)}
    return index, [row for row in rows if row[0] is not None]


@pytest.fixture(scope="module")
def real_clean():
    return clean_data(load_raw_data(DATA_PATH))


@needs_dataset
def test_real_dataset_has_all_required_columns():
    assert set(REQUIRED_COLUMNS).issubset(load_raw_data(DATA_PATH).columns)


@needs_dataset
def test_real_dataset_cleaning_accounts_for_every_row(real_clean):
    _, report = real_clean
    removed = (
        report["rows_dropped_missing_order_id"]
        + report["rows_dropped_invalid_date"]
        + report["exact_duplicate_rows_removed"]
    )
    assert report["rows_out"] == report["rows_in"] - removed
    index, rows = read_with_openpyxl(DATA_PATH)
    assert report["rows_in"] == len(rows)


@needs_dataset
def test_real_dataset_kpis_match_independent_python_calculation(real_clean):
    cleaned, _ = real_clean
    index, rows = read_with_openpyxl(DATA_PATH)
    sold = [r for r in rows if r[index["Order_Status"]] in ("Delivered", "Shipped")]

    expected_sales = math.fsum(r[index["Total_Sales_INR"]] for r in sold)
    expected_profit = math.fsum(r[index["Profit_INR"]] for r in sold)
    expected_units = sum(r[index["Quantity"]] for r in sold)
    expected_recorded = math.fsum(r[index["Total_Sales_INR"]] for r in rows)
    all_ids = {r[index["Order_ID"]] for r in rows}
    sold_ids = {r[index["Order_ID"]] for r in sold}
    returned_ids = {r[index["Order_ID"]] for r in rows if r[index["Order_Status"]] == "Returned"}
    cancelled_ids = {r[index["Order_ID"]] for r in rows if r[index["Order_Status"]] == "Cancelled"}

    k = metrics.compute_kpis(cleaned)
    assert k["total_sales"] == pytest.approx(expected_sales, abs=0.01)
    assert k["total_profit"] == pytest.approx(expected_profit, abs=0.01)
    assert k["units_sold"] == expected_units
    assert k["recorded_sales"] == pytest.approx(expected_recorded, abs=0.01)
    assert k["total_orders"] == len(all_ids)
    assert k["sold_orders"] == len(sold_ids)
    assert k["avg_order_value"] == pytest.approx(expected_sales / len(sold_ids))
    assert k["profit_margin_pct"] == pytest.approx(expected_profit / expected_sales * 100)
    assert k["return_rate_pct"] == pytest.approx(len(returned_ids) / len(all_ids) * 100)
    assert k["cancellation_rate_pct"] == pytest.approx(len(cancelled_ids) / len(all_ids) * 100)
    assert k["recorded_sales"] == pytest.approx(k["total_sales"] + k["lost_sales"], abs=0.01)


@needs_dataset
def test_real_dataset_state_orders_and_monthly_sales_match_independent_calculation(real_clean):
    cleaned, _ = real_clean
    index, rows = read_with_openpyxl(DATA_PATH)

    state_ids = defaultdict(set)
    monthly_sales = defaultdict(float)
    for r in rows:
        state_ids[r[index["Ship_State"]]].add(r[index["Order_ID"]])
        if r[index["Order_Status"]] in ("Delivered", "Shipped"):
            monthly_sales[str(r[index["Order_Date"]])[:7]] += r[index["Total_Sales_INR"]]

    by_state = metrics.summarize_by(cleaned, "Ship_State").set_index("Ship_State")["Orders"].to_dict()
    assert by_state == {state: len(ids) for state, ids in state_ids.items()}

    monthly = metrics.monthly_trend(cleaned)
    shown = {row.Month.strftime("%Y-%m"): row.Sales for row in monthly.itertuples()}
    for month, expected in monthly_sales.items():
        assert shown[month] == pytest.approx(expected, abs=0.01)
    assert list(monthly["Month"]) == sorted(monthly["Month"])
